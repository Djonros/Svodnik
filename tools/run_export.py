"""Полный экспорт сборки так же, как из окна программы: analyze + generate (Excel),
затем печать строк ведомости из получившегося Excel."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))
import pythoncom
import kompas_export_final as kef
from openpyxl import load_workbook

pythoncom.CoInitialize()
exp = kef.KompasExportFinal()
try:
    ok = exp.analyze(sys.argv[1])
finally:
    exp.close_all_docs()
print("analyze:", ok)
res = exp.generate({"excel"})
print("Проблемы:", exp.problems)
print("Расхождения:", exp.quantity_warnings)
print("Создано спецификаций:", exp.generated_specs)
ws = load_workbook(res["excel"]).active
for row in ws.iter_rows(values_only=True):
    if any(v not in (None, "") for v in row):
        print(" | ".join("" if v is None else str(v) for v in row))
